"""End-to-end tests: complete workflows through the real ingestion pipeline, real
BM25 + Chroma indexes and real page rendering, with deterministic fake models."""
from __future__ import annotations

import shutil
from pathlib import Path

import pytest

from course_assistant.answering import NOT_FOUND, answer_question
from course_assistant.app import build_app, format_answer
from course_assistant.clients import ServiceError
from course_assistant.documents import find_soffice
from course_assistant.quiz import generate_quiz
from course_assistant.store import Library
from tests.conftest import CAPTIONS, make_deck
from tests.fakes import FakeClients


# ----------------------------------------------------------------- managing materials
def test_add_duplicate_replace_and_remove(lib, deck, tmp_path):
    r = lib.add_file(deck)
    assert r["status"] == "added" and len(lib.list_docs()) == 1
    n_chunks, n_images = lib.text_col.count(), lib.image_col.count()
    assert n_chunks > 0 and n_images == 5

    # same file again (even under another path) -> no duplicate
    copy = tmp_path / "copy" / deck.name
    copy.parent.mkdir()
    shutil.copy(deck, copy)
    assert lib.add_file(copy)["status"] == "duplicate"
    assert len(lib.list_docs()) == 1 and lib.text_col.count() == n_chunks and lib.image_col.count() == n_images

    # same name, new content -> replaces the old version
    (tmp_path / "v2").mkdir()
    newer = make_deck(tmp_path / "v2" / deck.name, variant=" (updated)")
    r2 = lib.add_file(newer)
    assert r2["status"] == "replaced" and len(lib.list_docs()) == 1
    assert lib.image_col.count() == n_images  # old pages gone, new pages in

    # remove -> every index is empty, answers no longer use it
    doc_id = lib.list_docs()[0]["doc_id"]
    assert lib.remove(doc_id)
    assert lib.list_docs() == [] and lib.text_col.count() == 0 and lib.image_col.count() == 0 and lib.chunks == []
    assert not (lib.docs_dir / doc_id).exists()
    res = answer_question(lib, "What percent is the final project?")
    assert res["found"] is False and res["sources"] == []


def test_library_persists_across_restarts(settings, clients, deck):
    Library(settings, clients).add_file(deck)
    again = Library(settings, clients)
    assert len(again.list_docs()) == 1 and len(again.chunks) > 0
    assert again.keyword_search("final project", 3)[0]["page"] == 5


# ----------------------------------------------------------------- answering
def test_text_question_cites_correct_page(lib, deck):
    lib.add_file(deck)
    res = answer_question(lib, "What percentage of the grade is the final project?")
    assert res["found"] and "45%" in res["answer"]
    s = res["sources"][0]
    assert s["doc_name"] == deck.name and s["page"] == 5 and s["verified"] and Path(s["image"]).exists()
    assert res["support"]["verdict"] == "supported"
    md, src_md, gallery, details = format_answer(res)
    assert "Slide 5" in src_md and gallery and "Timings" in details


def test_visual_question_retrieves_and_shows_the_meme_slide(lib, deck):
    lib.add_file(deck)
    res = answer_question(lib, "Find the meme about vibe coding on Prod. What does the image say?")
    assert res["found"]
    s = res["sources"][0]
    assert s["page"] == 2 and s["evidence_type"] == "image" and Path(s["image"]).exists()
    assert "one does not simply" in res["answer"].lower()
    # the visual index (not just keyword) found it
    top = res["retrieved"][0]
    assert top["page"] == 2 and "visual" in top["ranks"]


def test_chart_question_uses_visual_evidence(lib, deck):
    lib.add_file(deck)
    res = answer_question(lib, "What does the Stack Overflow monthly questions chart show after ChatGPT launch?")
    assert res["found"] and res["sources"][0]["page"] == 3


def test_missing_information_is_acknowledged(lib, deck):
    lib.add_file(deck)
    res = answer_question(lib, "What is the professor's office phone number?")
    assert res["found"] is False and res["sources"] == []


def test_invented_citations_are_removed(lib, deck, clients):
    lib.add_file(deck)
    clients.llm.bad_citation = True
    res = answer_question(lib, "What percentage of the grade is the final project?")
    assert res["found"] is False and res["sources"] == [] and res["answer"] == NOT_FOUND


def test_invalid_json_is_repaired_with_one_retry(lib, deck, clients):
    lib.add_file(deck)
    clients.llm.garbage_first = True
    res = answer_question(lib, "What percentage of the grade is the final project?", support=False)
    assert res["found"] and clients.llm.calls == 2


def test_document_filter_limits_search(lib, deck, tmp_path):
    lib.add_file(deck)
    other = tmp_path / "notes.md"
    other.write_text("# Office hours\nOffice hours are Tuesdays and Fridays 11am to 12pm via Zoom.\n")
    r = lib.add_file(other)
    res = answer_question(lib, "When are office hours?", doc_ids=[r["doc_id"]])
    assert res["found"] and all(s["doc_id"] == r["doc_id"] for s in res["sources"])
    res2 = answer_question(lib, "When are office hours?", doc_ids=[lib.list_docs()[1]["doc_id"]
                                                                   if lib.list_docs()[0]["doc_id"] == r["doc_id"]
                                                                   else lib.list_docs()[0]["doc_id"]])
    assert all(s["doc_id"] != r["doc_id"] for s in res2["sources"])


def test_reranker_toggle(lib, deck, clients):
    lib.add_file(deck)
    on = answer_question(lib, "What does reranking do in hybrid RAG?", use_reranker=True, support=False)
    off = answer_question(lib, "What does reranking do in hybrid RAG?", use_reranker=False, support=False)
    assert on["reranked"] and not off["reranked"] and clients.reranker.calls == 1
    assert on["sources"][0]["page"] == 4 and off["sources"][0]["page"] == 4


# ----------------------------------------------------------------- unavailable services
def test_ingestion_rolls_back_when_embedding_service_is_down(lib, deck, clients):
    clients.visual_embedder.down = True
    with pytest.raises(ServiceError):
        lib.add_file(deck)
    assert lib.list_docs() == [] and lib.text_col.count() == 0 and lib.image_col.count() == 0
    assert list(lib.docs_dir.iterdir()) == []
    clients.visual_embedder.down = False
    assert lib.add_file(deck)["status"] == "added"


def test_llm_down_still_shows_sources(lib, deck, clients):
    lib.add_file(deck)
    clients.llm.down = True
    res = answer_question(lib, "What percentage of the grade is the final project?")
    assert res.get("llm_error") and res["sources"] and "unavailable" in res["answer"]


def test_retrievers_degrade_gracefully(lib, deck, clients):
    lib.add_file(deck)
    clients.reranker.down = True
    clients.text_embedder.down = True
    res = answer_question(lib, "What percentage of the grade is the final project?")
    assert res["found"] and any("reranker unavailable" in w for w in res["warnings"])
    assert any("text search unavailable" in w for w in res["warnings"])


def test_ocr_for_image_files(lib, tmp_path):
    from PIL import Image
    p = tmp_path / "handout.png"
    Image.new("RGB", (400, 300), "white").save(p)
    r = lib.add_file(p)
    assert r["status"] == "added" and lib.get_doc(r["doc_id"])["pages"][0]["ocr"]
    assert lib.keyword_search("prompt caching", 3)


# ----------------------------------------------------------------- quizzes
def test_quiz_generation_scoring_and_hidden_key(lib, deck):
    r = lib.add_file(deck)
    q = generate_quiz(lib, [r["doc_id"]], topic="", n=4, seed=7)
    assert len(q.questions) == 4
    for i, item in enumerate(q.questions):
        pub = q.public_question(i)
        assert set(pub) == {"question", "options"} and len(pub["options"]) == 4
        assert item["source"]["doc_name"] == deck.name and item["source"]["quote"]
        # the answer key points to the option text that is quoted in the source
        assert item["options"][item["correct_option"]] == item["source"]["quote"]
    key = [it["correct_option"] for it in q.questions]
    q.submit(0, key[0])
    q.submit(1, (key[1] + 1) % 4)
    assert q.score()["correct"] == 1
    assert [it["correct_option"] for it in q.questions] == key  # key unchanged by answering


def test_quiz_topic_filter_and_errors(lib, deck):
    r = lib.add_file(deck)
    q = generate_quiz(lib, [r["doc_id"]], topic="hybrid RAG reranking", n=2, seed=1)
    assert any(it["source"]["page"] == 4 for it in q.questions)
    with pytest.raises(ValueError):
        generate_quiz(lib, [], "", 3)


# ----------------------------------------------------------------- PowerPoint via LibreOffice
@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed")
def test_pptx_upload_is_converted_to_slide_images(lib, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    for title, body in (("Week 9 Fine-Tuning", "LoRA adds low-rank adapters."), ("QLoRA", "QLoRA quantizes the base model to 4-bit.")):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
        s.notes_slide.notes_text_frame.text = f"Speaker note for {title}"
    path = tmp_path / "Week 9.pptx"
    prs.save(path)
    r = lib.add_file(path)
    doc = lib.get_doc(r["doc_id"])
    assert doc["kind"] == "slides" and doc["n_pages"] == 2
    assert all(Path(p["image"]).exists() for p in doc["pages"])
    assert "Speaker note for QLoRA" in lib.page_text(r["doc_id"], 2)
    res = answer_question(lib, "What does QLoRA quantize the base model to?")
    assert res["found"] and res["sources"][0]["label"] == "Slide 2"


# ----------------------------------------------------------------- UI wiring
def test_gradio_app_builds(settings, deck):
    demo, lib = build_app(settings, FakeClients(CAPTIONS))
    assert demo is not None and lib.list_docs() == []


# ----------------------------------------------------------------- evaluation harness
def test_eval_harness_scores_and_summarizes(lib, deck):
    from eval.run_eval import evaluate, keyword_ok, summarize
    lib.add_file(deck)
    qs = [{"id": "A", "type": "text", "answerable": True, "question": "What percentage of the grade is the final project?",
           "keywords": [["45"]], "expected_sources": [{"doc": "Week 2", "pages": [5]}]},
          {"id": "B", "type": "missing", "answerable": False, "question": "What is the professor's phone number?",
           "keywords": [], "expected_sources": []}]
    rows = {c: evaluate(lib, qs, c, 3) for c in ("rerank_on", "rerank_off")}
    assert all(r["auto_correct"] and r["cited_expected_source"] for rs in rows.values() for r in rs)
    md = summarize(rows)
    assert "rerank_on" in md and "2/2" in md
    assert keyword_ok("It is forty-five percent, i.e. 45%", [["45"]]) and not keyword_ok("unknown", [["45"]])
